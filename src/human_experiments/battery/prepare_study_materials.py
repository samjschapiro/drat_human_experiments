"""Build reproducible, validated DRAT assignments from the approved workbooks.

Usage: python prepare_study_materials.py --materials-dir PATH

The input workbooks are supplied by the study team and are deliberately not
copied into the public web directory. The generated JavaScript contains only
stimuli and assignments, never participant data or scoring keys.
"""

from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
import random

from openpyxl import load_workbook


HERE = Path(__file__).parent
RELATIONS = ("causal", "constitutive", "categorical", "unrelated")
SEED = 20260929


def read_materials(materials_dir: Path) -> tuple[list[dict], list[dict]]:
    triplet_sheet = load_workbook(
        materials_dir / "All_Triplets.xlsx", read_only=True, data_only=True
    ).active
    triplets = [
        {
            "triplet_id": int(row[0]),
            "relation": str(row[4]).strip().lower(),
            "anchors": [str(word).strip().lower() for word in row[1:4]],
        }
        for row in triplet_sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    ]
    by_id = {row["triplet_id"]: row for row in triplets}
    if len(triplets) != 75 or len(by_id) != 75:
        raise ValueError("Expected 75 uniquely numbered triplets")
    expected = {"causal": 10, "constitutive": 10, "categorical": 16, "unrelated": 39}
    if Counter(row["relation"] for row in triplets) != expected:
        raise ValueError("Triplet relation counts differ from the approved materials")

    pair_sheet = load_workbook(
        materials_dir / "DRAT_k2_Anchor_Pairs.xlsx", read_only=True, data_only=True
    ).active
    pairs = []
    for index, row in enumerate(pair_sheet.iter_rows(min_row=2, values_only=True), 1):
        if row[0] is None:
            continue
        triplet_id = int(row[0])
        relation = str(row[1]).strip().lower()
        anchors = [str(row[3]).strip().lower(), str(row[4]).strip().lower()]
        if triplet_id not in by_id or relation != by_id[triplet_id]["relation"]:
            raise ValueError(f"Pair row {index} has an invalid source triplet/relation")
        if len(set(anchors)) != 2 or not set(anchors) <= set(by_id[triplet_id]["anchors"]):
            raise ValueError(f"Pair row {index} is not a distinct pair from its triplet")
        selection = str(row[5]).strip().lower()
        if (relation == "unrelated") != selection.startswith("eligible"):
            raise ValueError(f"Pair row {index} has the wrong selection rule")
        pairs.append({
            "pair_id": f"pair_{index:03d}",
            "triplet_id": triplet_id,
            "relation": relation,
            "anchors": anchors,
            "selection": selection,
        })
    if len(pairs) != 153:
        raise ValueError("Expected 153 two-word pair rows")
    if Counter(row["relation"] == "unrelated" for row in pairs) != {True: 117, False: 36}:
        raise ValueError("Expected 117 unrelated and 36 related two-word pairs")
    pair_counts = Counter(row["triplet_id"] for row in pairs)
    for triplet in triplets:
        want = 3 if triplet["relation"] == "unrelated" else 1
        if pair_counts[triplet["triplet_id"]] != want:
            raise ValueError(f"Triplet {triplet['triplet_id']} has the wrong number of pair options")
    return triplets, pairs


def generate_assignments(triplets: list[dict], pairs: list[dict], count: int) -> list[dict]:
    pools = {}
    for relation in RELATIONS:
        pools[(3, relation)] = [row for row in triplets if row["relation"] == relation]
        pools[(2, relation)] = [row for row in pairs if row["relation"] == relation]
    cells = [(size, relation) for relation in RELATIONS for size in (2, 3)]
    source_usage = Counter()
    pair_usage = Counter()
    assignments = []

    for slot in range(count):
        rng = random.Random(SEED + slot)

        def search(position: int, used_words: set[str], used_sources: set[int],
                   chosen: list[dict]) -> list[dict] | None:
            if position == len(cells):
                return chosen
            size, relation = cells[position]
            candidates = pools[(size, relation)].copy()
            rng.shuffle(candidates)
            candidates.sort(key=lambda item: (
                source_usage[(size, relation, item["triplet_id"])],
                pair_usage[item["pair_id"]] if size == 2 else 0,
            ))
            for item in candidates:
                words = set(item["anchors"])
                if item["triplet_id"] in used_sources or words & used_words:
                    continue
                result = search(position + 1, used_words | words,
                                used_sources | {item["triplet_id"]},
                                chosen + [{"anchor_size": size, **item}])
                if result is not None:
                    return result
            return None

        blocks = search(0, set(), set(), [])
        if blocks is None:
            raise ValueError(f"Could not construct a valid eight-cell assignment for slot {slot}")
        for block in blocks:
            source_usage[(block["anchor_size"], block["relation"], block["triplet_id"])] += 1
            if block["anchor_size"] == 2:
                pair_usage[block["pair_id"]] += 1
                if block["relation"] == "unrelated" and rng.randrange(2):
                    block["anchors"] = list(reversed(block["anchors"]))
        rng.shuffle(blocks)
        for position, block in enumerate(blocks):
            block["presentation_index"] = position
            block["set_id"] = f"slot_{slot:03d}_{block['relation']}_k{block['anchor_size']}"
        validate_assignment(blocks)
        assignments.append({"slot": slot, "seed": SEED + slot, "blocks": blocks})

    # If a source can be chosen, usage should not drift far from its peers.
    for size, relation in cells:
        values = [source_usage[(size, relation, row["triplet_id"])]
                  for row in pools[(size, relation)]]
        if max(values) - min(values) > 3:
            raise ValueError(f"Unbalanced {relation} k={size} assignments: {min(values)}–{max(values)}")
    return assignments


def validate_assignment(blocks: list[dict]) -> None:
    if len(blocks) != 8:
        raise ValueError("DRAT assignment must have eight blocks")
    if Counter((b["anchor_size"], b["relation"]) for b in blocks) != {
        (size, relation): 1 for relation in RELATIONS for size in (2, 3)
    }:
        raise ValueError("DRAT assignment must cover every k × relation cell once")
    words = [word for block in blocks for word in block["anchors"]]
    if len(words) != 20 or len(set(words)) != 20:
        raise ValueError("DRAT assignment must contain twenty distinct anchor words")
    if len({block["triplet_id"] for block in blocks}) != 8:
        raise ValueError("A pair and its source triplet cannot appear together")
    if sorted(block["presentation_index"] for block in blocks) != list(range(8)):
        raise ValueError("DRAT presentation positions must be unique")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--materials-dir", type=Path, required=True)
    parser.add_argument("--count", type=int, default=260)
    parser.add_argument("--output", type=Path, default=HERE / "js" / "study-materials.js")
    args = parser.parse_args()
    if args.count < 1:
        raise ValueError("Assignment count must be positive")
    triplets, pairs = read_materials(args.materials_dir)
    assignments = generate_assignments(triplets, pairs, args.count)
    payload = {
        "protocol_version": "drat-human-2026-v1",
        "triplets": triplets,
        "pairs": pairs,
        "assignments": assignments,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        "// Generated from approved study workbooks; do not edit by hand.\n"
        + "window.STUDY_MATERIALS = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {len(assignments)} validated assignments to {args.output}")


if __name__ == "__main__":
    main()
